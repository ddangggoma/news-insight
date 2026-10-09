"""Report export (plan 16 #11): one plain document model — sections of paragraphs and
bullets citing cards — rendered to Markdown, Word and PowerPoint, so a briefing, a weekly or
monthly briefing and a topic dossier leave the site in the form a report needs."""

import io
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from typing import TYPE_CHECKING

from docx import Document as WordDocument
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches
from pptx.util import Pt as SlidePt

if TYPE_CHECKING:
    from fastapi.responses import Response


@dataclass
class Source:
    item_id: int
    title: str
    url: str
    source: str


@dataclass
class Block:
    text: str
    refs: list[int] = field(default_factory=list)  # item ids
    bullet: bool = False
    heading: bool = False  # a sub-heading inside the section


@dataclass
class Section:
    title: str
    blocks: list[Block] = field(default_factory=list)


@dataclass
class Doc:
    title: str
    subtitle: str
    sections: list[Section]
    sources: dict[int, Source]  # item id -> source; numbered in order of first citation

    def numbers(self) -> dict[int, int]:
        order: dict[int, int] = {}
        for section in self.sections:
            for block in section.blocks:
                for ref in block.refs:
                    if ref in self.sources and ref not in order:
                        order[ref] = len(order) + 1
        return order

    def cite(self, block: Block, numbers: dict[int, int]) -> str:
        marks = "".join(f"[{numbers[r]}]" for r in block.refs if r in numbers)
        return f"{block.text} {marks}".rstrip()


def markdown(doc: Doc) -> str:
    numbers = doc.numbers()
    lines = [f"# {doc.title}", "", f"_{doc.subtitle}_", ""]
    for section in doc.sections:
        lines += [f"## {section.title}", ""]
        for block in section.blocks:
            text = doc.cite(block, numbers)
            if block.heading:
                lines += [f"### {text}", ""]
            elif block.bullet:
                lines.append(f"- {text}")
            else:
                lines += [text, ""]
        lines.append("")
    if numbers:
        lines += ["## 근거", ""]
        for item_id, n in sorted(numbers.items(), key=lambda kv: kv[1]):
            src = doc.sources[item_id]
            lines.append(f"{n}. [{src.title}]({src.url}) — {src.source}")
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip() + "\n"


def word(doc: Doc) -> bytes:
    numbers = doc.numbers()
    out = WordDocument()
    style = out.styles["Normal"]
    style.font.size = Pt(10.5)
    out.add_heading(doc.title, level=0)
    out.add_paragraph().add_run(doc.subtitle).italic = True
    for section in doc.sections:
        out.add_heading(section.title, level=1)
        for block in section.blocks:
            text = doc.cite(block, numbers)
            if block.heading:
                out.add_heading(text, level=2)
            elif block.bullet:
                out.add_paragraph(text, style="List Bullet")
            else:
                out.add_paragraph(text)
    if numbers:
        out.add_heading("근거", level=1)
        for item_id, n in sorted(numbers.items(), key=lambda kv: kv[1]):
            src = doc.sources[item_id]
            out.add_paragraph(f"[{n}] {src.title} — {src.source}\n{src.url}")
    buffer = io.BytesIO()
    out.save(buffer)
    return buffer.getvalue()


SLIDE_LINES = 7  # bullets per slide before a section continues on the next


def _slide(deck: object, title: str, lines: list[str]) -> None:
    layout = deck.slide_layouts[1]  # type: ignore[attr-defined]
    slide = deck.slides.add_slide(layout)  # type: ignore[attr-defined]
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.clear()
    for index, line in enumerate(lines):
        paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
        paragraph.text = line
        paragraph.font.size = SlidePt(14 if len(line) < 120 else 12)


def slides(doc: Doc) -> bytes:
    numbers = doc.numbers()
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = doc.title
    cover.placeholders[1].text = doc.subtitle
    for section in doc.sections:
        lines = [
            ("■ " if block.heading else "") + doc.cite(block, numbers)[:400]
            for block in section.blocks
            if block.text
        ]
        for start in range(0, max(len(lines), 1), SLIDE_LINES):
            part = lines[start : start + SLIDE_LINES]
            suffix = f" ({start // SLIDE_LINES + 1})" if len(lines) > SLIDE_LINES else ""
            _slide(deck, section.title + suffix, part or ["–"])
    if numbers:
        refs = [
            f"[{n}] {doc.sources[i].title[:90]} — {doc.sources[i].source}"
            for i, n in sorted(numbers.items(), key=lambda kv: kv[1])
        ]
        for start in range(0, len(refs), 10):
            _slide(deck, "근거", refs[start : start + 10])
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


FORMATS: dict[str, tuple[str, Callable[[Doc], str | bytes]]] = {
    "md": ("text/markdown; charset=utf-8", markdown),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", word),
    "pptx": ("application/vnd.openxmlformats-officedocument.presentationml.presentation", slides),
}


def render(doc: Doc, fmt: str) -> tuple[bytes, str]:
    content_type, renderer = FORMATS[fmt]
    body = renderer(doc)
    return (body.encode() if isinstance(body, str) else body), content_type


def download(doc: Doc, fmt: str, stem: str) -> "Response":
    """The rendered document as an attachment (ASCII file name plus the UTF-8 one)."""
    from urllib.parse import quote

    from fastapi.responses import Response

    body, content_type = render(doc, fmt)
    ascii_name = re.sub(r"[^A-Za-z0-9._-]+", "-", stem).strip("-") or "report"
    disposition = (
        f"attachment; filename=\"{ascii_name}.{fmt}\"; filename*=UTF-8''{quote(f'{stem}.{fmt}')}"
    )
    return Response(body, media_type=content_type, headers={"Content-Disposition": disposition})
