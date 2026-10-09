import io

import pytest
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from pptx import Presentation
from sqlalchemy.orm import Session

from news_insight.export.document import Block, Doc, Section, Source, markdown, slides, word
from tests.public.test_briefings_api import published
from tests.public.test_dossiers_api import client, headers  # noqa: F401 - fixtures

DOC = Doc(
    title="온디바이스 AI",
    subtitle="테스트",
    sections=[
        Section("요약", [Block("첫 줄", bullet=True), Block("본문", refs=[7, 3])]),
        Section("흐름", [Block("소제목", heading=True), Block("둘째", refs=[3, 99])]),
    ],
    sources={
        3: Source(3, "Card three", "https://ex.com/3", "A"),
        7: Source(7, "Card seven", "https://ex.com/7", "B"),
    },
)


def test_citations_are_numbered_in_order_of_first_use() -> None:
    text = markdown(DOC)
    assert "본문 [1][2]" in text and "둘째 [2]" in text  # unknown id 99 is dropped
    assert "1. [Card seven](https://ex.com/7) — B" in text and "### 소제목" in text


def test_word_and_slides_open_with_their_libraries() -> None:
    document = WordDocument(io.BytesIO(word(DOC)))
    assert any("본문 [1][2]" in p.text for p in document.paragraphs)
    deck = Presentation(io.BytesIO(slides(DOC)))
    titles = [s.shapes.title.text for s in deck.slides]
    assert titles == ["온디바이스 AI", "요약", "흐름", "근거"]


@pytest.mark.db
def test_briefing_exports_in_each_format(
    db_session: Session, public_client: TestClient, public_headers: dict[str, str]
) -> None:
    briefing = published(db_session)
    day = briefing.briefing_date.isoformat()
    for fmt, kind in (
        ("md", "text/markdown"),
        ("docx", "wordprocessingml"),
        ("pptx", "presentationml"),
    ):
        response = public_client.get(
            f"/api/public/export/briefing/{day}", params={"format": fmt}, headers=public_headers
        )
        assert response.status_code == 200, response.text
        assert kind in response.headers["content-type"]
        assert f'filename="briefing-{day}.{fmt}"' in response.headers["content-disposition"]
    text = public_client.get(f"/api/public/export/briefing/{day}", headers=public_headers).text
    assert text.startswith("# ") and "## 핵심 인사이트" in text and "## 근거" in text
    missing = public_client.get("/api/public/export/briefing/2001-01-01", headers=public_headers)
    assert missing.status_code == 404


@pytest.mark.db
def test_dossier_exports_with_its_hypotheses(
    client: TestClient,  # noqa: F811
    headers: dict[str, str],  # noqa: F811
) -> None:
    dossier = client.post(
        "/api/public/dossiers",
        json={"title": "에이전트 OS", "keywords": ["agent", "에이전트"]},
        headers=headers,
    ).json()
    base = f"/api/public/dossiers/{dossier['id']}"
    client.post(f"{base}/hypotheses", json={"text": "에이전트 OS가 기본이 된다"}, headers=headers)
    text = client.get(f"{base}/export", headers=headers).text
    assert (
        text.startswith("# 에이전트 OS")
        and "## 가설과 근거" in text
        and "에이전트 OS가 기본이 된다" in text
    )
    deck = client.get(f"{base}/export", params={"format": "pptx"}, headers=headers)
    assert deck.status_code == 200 and Presentation(io.BytesIO(deck.content)).slides
